"""Genera Practica_kLa.docx con la plantilla y la presentación del taller.

Uso:
    .venv/bin/python practica/generar_word.py

1. pandoc convierte el Markdown con sus estilos estándar (la plantilla como
   documento de referencia rompe las tablas y ecuaciones al renderizar).
2. python-docx le pasa lo de la plantilla: tamaño de página y márgenes, encabezado
   con logos, pie de página, colores y tamaños de títulos, Arial 11 en el cuerpo,
   las líneas de autor y curso bajo el título, y tablas con encabezado morado.
"""
import copy
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

AQUI = Path(__file__).parent
MD = AQUI / "Practica_kLa.md"
SALIDA = AQUI / "Practica_kLa.docx"
PLANTILLA = AQUI.parent / "plantilla_documento.docx"

AUTOR = "Octavio T. Ramírez · Emiliano Balderas · Instituto de Biotecnología, UNAM"

MORADO = RGBColor(0x4E, 0x28, 0x91)
MAGENTA = RGBColor(0xBD, 0x28, 0x74)
TINTA_SUAVE = RGBColor(0x4A, 0x3A, 0x68)
GRIS_MORADO = RGBColor(0x7A, 0x6B, 0x95)
GRIS_BORDE = "B7AFC6"

# (tamaño pt, color, antes pt, después pt) tomados de la plantilla
TITULOS = {
    "Title": (22, MORADO, 0, 4),
    "Heading1": (15, MORADO, 16, 6),
    "Heading2": (12, MAGENTA, 12, 5),
    "Heading3": (11, MORADO, 10, 4),
}

# Orden de los hijos que exige el esquema de Word
ORDEN_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
ORDEN_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
              "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
ORDEN_TRPR = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
              "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"]
ORDEN_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku",
             "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
             "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
             "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
             "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]


def poner(padre, hijo, orden):
    """Inserta 'hijo' en 'padre' respetando el orden del esquema (reemplaza si ya existe)."""
    nombre = hijo.tag.split("}")[1]
    for viejo in padre.findall(qn(f"w:{nombre}")):
        padre.remove(viejo)
    pos = orden.index(nombre)
    for i, el in enumerate(padre):
        n = el.tag.split("}")[1]
        if n in orden and orden.index(n) > pos:
            el.addprevious(hijo)
            return
    padre.append(hijo)


def elemento(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


# ---------------------------------------------------------------- página, encabezado y pie

def copiar_parte(destino, origen):
    """Copia el contenido de un encabezado/pie de la plantilla, incluidas sus imágenes."""
    destino.is_linked_to_previous = False
    parte_d, parte_o = destino.part, origen.part
    remap = {}
    for rid, rel in parte_o.rels.items():
        if rel.reltype == RT.IMAGE:
            remap[rid] = parte_d.relate_to(rel.target_part, RT.IMAGE)
    el_d = destino._element
    for hijo in list(el_d):
        el_d.remove(hijo)
    for hijo in origen._element:
        nuevo = copy.deepcopy(hijo)
        for nodo in nuevo.iter():
            for att in (qn("r:embed"), qn("r:id"), qn("r:link")):
                if nodo.get(att) in remap:
                    nodo.set(att, remap[nodo.get(att)])
        el_d.append(nuevo)


def pagina(doc, plantilla):
    so, sd = plantilla.sections[0], doc.sections[0]
    for att in ("page_width", "page_height", "left_margin", "right_margin", "top_margin",
                "bottom_margin", "header_distance", "footer_distance"):
        setattr(sd, att, getattr(so, att))
    copiar_parte(sd.header, so.header)
    copiar_parte(sd.footer, so.footer)


# ---------------------------------------------------------------- estilos

def fuente_arial(estilo, tamano=None):
    estilo.font.name = "Arial"
    rfonts = estilo.element.rPr.find(qn("w:rFonts"))
    for att in list(rfonts.attrib):
        if "Theme" in att:
            del rfonts.attrib[att]
    rfonts.set(qn("w:cs"), "Arial")
    rfonts.set(qn("w:eastAsia"), "Arial")
    if tamano:
        estilo.font.size = Pt(tamano)


def estilos(doc):
    # Valores por omisión: Arial 11
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rfonts = rpr.find(qn("w:rFonts"))
    for att in list(rfonts.attrib):
        del rfonts.attrib[att]
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{k}"), "Arial")
    for st in doc.styles:
        if st.type not in (1, 2) or st.name in ("Source Code", "Verbatim Char"):
            continue
        fuente_arial(st)
    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        st = doc.styles[nombre]
        fuente_arial(st, 11)
        st.font.color.rgb = RGBColor(0x1F, 0x1A, 0x2B)
    for nombre in ("Body Text", "First Paragraph"):
        pf = doc.styles[nombre].paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(6), 1.15
    pf = doc.styles["Compact"].paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(1), Pt(1), 1.0

    for sid, (tam, color, antes, despues) in TITULOS.items():
        st = next(s for s in doc.styles if s.style_id == sid)
        fuente_arial(st, tam)
        st.font.bold, st.font.italic = True, False
        st.font.color.rgb = color
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(antes), Pt(despues)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Notas (> en Markdown): sangría y raya magenta a la izquierda
    bt = doc.styles["Block Text"]
    bt.font.italic = False
    bt.font.color.rgb = TINTA_SUAVE
    bt.paragraph_format.left_indent = Cm(0.5)
    bt.paragraph_format.space_before, bt.paragraph_format.space_after = Pt(4), Pt(4)
    ppr = bt.element.get_or_add_pPr()
    bdr = elemento("w:pBdr")
    bdr.append(elemento("w:left", val="single", sz=18, space=8, color="BD2874"))
    poner(ppr, bdr, ORDEN_PPR)

    for nombre in ("Image Caption", "Table Caption"):
        st = doc.styles[nombre]
        st.font.size, st.font.italic, st.font.color.rgb = Pt(9), True, GRIS_MORADO


def lineas_bajo_titulo(doc, plantilla):
    """Copia de la plantilla las dos líneas bajo el título (autor y curso) con su formato."""
    titulo = next(p for p in doc.paragraphs if p.style.style_id == "Title")
    autor = copy.deepcopy(plantilla.paragraphs[1]._p)
    curso = copy.deepcopy(plantilla.paragraphs[2]._p)
    runs = autor.findall(qn("w:r"))
    runs[0].find(qn("w:t")).text = AUTOR
    for r in runs[1:]:
        autor.remove(r)
    titulo._p.addnext(curso)
    titulo._p.addnext(autor)


# ---------------------------------------------------------------- tablas y figuras

def tablas(doc):
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        poner(tblpr, elemento("w:tblW", type="pct", w=5000), ORDEN_TBLPR)
        bordes = elemento("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            bordes.append(elemento(f"w:{lado}", val="single", sz=4, space=0, color=GRIS_BORDE))
        poner(tblpr, bordes, ORDEN_TBLPR)
        margenes = elemento("w:tblCellMar")
        for lado, w in (("top", 50), ("left", 100), ("bottom", 50), ("right", 100)):
            margenes.append(elemento(f"w:{lado}", w=w, type="dxa"))
        poner(tblpr, margenes, ORDEN_TBLPR)

        # Encabezado morado con texto blanco
        for celda in t.rows[0].cells:
            poner(celda._tc.get_or_add_tcPr(), elemento("w:shd", val="clear", color="auto", fill="4E2891"),
                  ORDEN_TCPR)
            for p in celda.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Tablas para llenar a mano: renglones más altos
        celdas = [c for row in t.rows[1:] for c in row.cells]
        if celdas and sum(not c.text.strip() for c in celdas) / len(celdas) > 0.3:
            for row in t.rows[1:]:
                row.height = Cm(0.8)

        # Renglones sin partirse entre páginas; el encabezado se repite en cada página
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            poner(trpr, elemento("w:cantSplit"), ORDEN_TRPR)
            if i == 0:
                poner(trpr, elemento("w:tblHeader"), ORDEN_TRPR)

        # Aire entre la tabla y el párrafo que sigue
        sig = t._tbl.getnext()
        if sig is not None and sig.tag == qn("w:p"):
            Paragraph(sig, t._parent).paragraph_format.space_before = Pt(8)


def figuras(doc, ancho_cm=14):
    for forma in doc.inline_shapes:
        proporcion = forma.height / forma.width
        ancho = 12.5 if proporcion > 1.2 else ancho_cm
        forma.width = Cm(ancho)
        forma.height = int(Cm(ancho) * proporcion)


def saltos_de_pagina(doc, inicios=("6. Registro experimental", "Anexo para el instructor")):
    """La hoja de datos y el anexo empiezan en página nueva (se imprimen o quitan por separado)."""
    for p in doc.paragraphs:
        if p.style.style_id == "Heading1" and p.text.startswith(inicios):
            p.paragraph_format.page_break_before = True


def presentacion(doc):
    """Conserva el texto justificado del Word editado y centra figuras y ecuaciones."""
    for p in doc.paragraphs:
        if p.style.style_id == "Title":
            continue
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if p._p.xpath(".//w:drawing") or p._p.xpath(".//m:oMathPara"):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = bool(p._p.xpath(".//w:drawing"))
        if p.style.name == "Image Caption":
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = False
    # Mantener la introducción de cada ecuación junto a la fórmula.
    for anterior, siguiente in zip(doc.paragraphs, doc.paragraphs[1:]):
        if siguiente._p.xpath(".//m:oMathPara"):
            anterior.paragraph_format.keep_with_next = True
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def generar(md, salida):
    subprocess.run(["pandoc", str(md), "-f", "markdown+tex_math_dollars", "-o", str(salida),
                    f"--resource-path={AQUI}:{AQUI.parent}"], check=True)
    doc = Document(salida)
    plantilla = Document(PLANTILLA)
    pagina(doc, plantilla)
    estilos(doc)
    lineas_bajo_titulo(doc, plantilla)
    tablas(doc)
    figuras(doc)
    saltos_de_pagina(doc)
    presentacion(doc)
    doc.save(salida)
    print(f"Listo: {salida}")


def main():
    import argparse
    parser = argparse.ArgumentParser(description="Genera la guía y el anexo del instructor.")
    parser.add_argument("--pdf", action="store_true", help="Exporta ambos PDF con LibreOffice.")
    args = parser.parse_args()
    documentos = [(MD, SALIDA), (AQUI / "Anexo_instructor.md", AQUI / "Anexo_instructor.docx")]
    for md, salida in documentos:
        generar(md, salida)
    if args.pdf:
        subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(AQUI),
                        *(str(salida) for _, salida in documentos)], check=True)


if __name__ == "__main__":
    main()
