"""Genera la plantilla de cálculo sin datos experimentales. Requiere openpyxl."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

BASE=Path(__file__).parent
wb=Workbook(); summary=wb.active;summary.title='Barrido'
wb.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True)
readme=wb.create_sheet('Instrucciones')
lines=[
'Plantilla de cálculo de kLa: reactor Sartorius 50 L, 30 L de agua, aire 15 L/min.',
'Una hoja por velocidad. No hay datos experimentales precargados.',
'Introducir tiempo en segundos desde el cambio de gas en A y oxígeno en B (filas 9 a 308).',
'Introducir meseta inicial en K2 y meseta final en K3, tras revisar el registro completo.',
'No sustituir automáticamente la meseta final por 100 %. La escala debe ser la misma.',
'Las columnas C y D calculan el recorrido y la transformación logarítmica.',
'E y F seleccionan Y y tiempo solo para un recorrido entre 0.20 y 0.80.',
'El ajuste tiene ordenada libre. Se muestran resultados a partir de tres puntos válidos.',
'Tres puntos son un mínimo de cálculo, no un criterio suficiente de calidad.',
'K5 expresa kLa en h⁻¹. Revisar pendiente positiva, linealidad, residuos y condiciones.',
'G muestra los residuos Y − (intercepto + pendiente × tiempo) de los puntos seleccionados.',
'Las gráficas de Y y residuos representan el tramo seleccionado; la curva de oxígeno usa todo el registro.',
'Campos azules: entrada. Conservar una copia por equipo y los registros originales.',
'Las fórmulas se calculan al abrir en Excel o LibreOffice. Las celdas vacías no representan kLa = 0.',
'La plantilla admite 300 muestras por curva. Si el registro es mayor, ampliar fórmulas, rangos y gráficas.',
]
for i,line in enumerate(lines,1):readme.cell(i,1,line)
readme.column_dimensions['A'].width=125
for row in readme: row[0].alignment=Alignment(wrap_text=True,vertical='center');readme.row_dimensions[row[0].row].height=32
summary.append(['RPM','kLa [h⁻¹]','R²','Puntos','Temperatura [°C]','Observaciones'])
for idx,rpm in enumerate([200,400,600,800,1000],2):
 ws=wb.create_sheet(str(rpm)+' rpm');name=ws.title
 ws['A1']=f'Reoxigenación a {rpm} rpm';ws['A2']='Volumen: 30 L de agua';ws['A3']='Aire: 15 L/min = 0.5 vvm'
 ws['A4']='Archivo / equipo:';ws['A5']='Temperatura [°C]:';ws['A6']='Intervalos de mesetas:'
 labels=['RPM','y₀ (entrada)','y∞ (entrada)','Pendiente [s⁻¹]','kLa [h⁻¹]','R²','Intercepto','Puntos de ajuste']
 for i,label in enumerate(labels,1):ws.cell(i,10,label)
 ws['K1']=rpm
 ws['K4']='=IF(AND(K8>=3,ISNUMBER(K2),ISNUMBER(K3),K3>K2),IFERROR(SLOPE(E9:E308,F9:F308),""),"")'
 ws['K5']='=IF(ISNUMBER(K4),K4*3600,"")'
 ws['K6']='=IF(ISNUMBER(K4),IFERROR(RSQ(E9:E308,F9:F308),""),"")'
 ws['K7']='=IF(ISNUMBER(K4),IFERROR(INTERCEPT(E9:E308,F9:F308),""),"")'
 ws['K8']='=COUNT(E9:E308)'
 ws.append([])
 for c,h in enumerate(['Tiempo [s]','Oxígeno [%]','Recorrido r','Y','Y seleccionado','t seleccionado','Residuo'],1):ws.cell(8,c,h)
 # J/K8 contain metadata and the selected count; data start at row 9.
 for r in range(9,309):
  ws.cell(r,3,f'=IF(AND(ISNUMBER(A{r}),ISNUMBER(B{r}),ISNUMBER($K$2),ISNUMBER($K$3),$K$3>$K$2),(B{r}-$K$2)/($K$3-$K$2),"")')
  ws.cell(r,4,f'=IF(ISNUMBER(C{r}),IF(B{r}<$K$3,LN(($K$3-$K$2)/($K$3-B{r})),""),"")')
  ws.cell(r,5,f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),C{r}>=0.2,C{r}<=0.8),D{r},"")')
  ws.cell(r,6,f'=IF(ISNUMBER(E{r}),A{r},"")')
  ws.cell(r,7,f'=IF(AND(ISNUMBER(E{r}),ISNUMBER($K$4)),E{r}-($K$7+$K$4*A{r}),"")')
  for c in [1,2]:ws.cell(r,c).fill=PatternFill('solid',fgColor='E4EFFA')
  for c in range(3,8):ws.cell(r,c).number_format='0.0000'
 for key in ['K2','K3','B4','B5','B6']:ws[key].fill=PatternFill('solid',fgColor='E4EFFA')
 for col in 'ABCDEFG':ws.column_dimensions[col].width=19
 ws.column_dimensions['J'].width=24;ws.column_dimensions['K'].width=18
 ws.freeze_panes='C9';ws.auto_filter.ref='A8:G308'
 for title,xcol,ycol,anchor,ytitle in [('Reoxigenación',1,2,'I11','Oxígeno [%]'),('Ajuste: tramo seleccionado',6,5,'I27','Y'),('Residuos',6,7,'I43','Residuo Y')]:
  chart=ScatterChart();chart.title=title;chart.x_axis.title='Tiempo [s]';chart.y_axis.title=ytitle;chart.width=20;chart.height=8
  series=Series(Reference(ws,min_col=ycol,min_row=9,max_row=308),Reference(ws,min_col=xcol,min_row=9,max_row=308),title=title)
  series.marker.symbol='circle';series.marker.size=3;series.graphicalProperties.line.noFill=True
  if ycol==5:
   from openpyxl.chart.trendline import Trendline
   series.trendline=Trendline(trendlineType='linear',dispEq=True,dispRSqr=True)
  chart.series.append(series);ws.add_chart(chart,anchor)
 summary.append([rpm,f'=IF(ISNUMBER(\'{name}\'!K5),\'{name}\'!K5,"")',f'=IF(ISNUMBER(\'{name}\'!K6),\'{name}\'!K6,"")',f"='{name}'!K8",None,None])
chart=ScatterChart();chart.title='Barrido de kLa a 0.5 vvm y 30 L';chart.x_axis.title='Agitación [rpm]';chart.y_axis.title='kLa [h⁻¹]';chart.width=23;chart.height=13
series=Series(Reference(summary,min_col=2,min_row=2,max_row=6),Reference(summary,min_col=1,min_row=2,max_row=6),title='kLa');series.marker.symbol='circle';series.graphicalProperties.line.noFill=True;chart.series.append(series);summary.add_chart(chart,'A10')
for col in 'ABCDEF':summary.column_dimensions[col].width=23
summary.column_dimensions['F'].width=45
for ws in wb:
 for row in [ws[1]] if ws.title in ['Barrido','Instrucciones'] else [ws[8]]:
  for cell in row:
   if cell.value is not None:cell.fill=PatternFill('solid',fgColor='4E2891');cell.font=Font(name='Arial',bold=True,color='FFFFFF')
 ws.sheet_view.showGridLines=False
wb.save(BASE/'Calculo_kLa.xlsx')
print('Plantilla generada:',BASE/'Calculo_kLa.xlsx')
